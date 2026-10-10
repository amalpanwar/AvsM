from __future__ import annotations

import argparse
import json
import math
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


ROOT = Path(__file__).resolve().parents[1]
HISTORY_FILE = ROOT / "epl-2025-GMTStandardTime.xlsx"
FIXTURE_FILE = ROOT / "epl-2026-GMTStandardTime.xlsx"
OUT_FILE = ROOT / "generated-data.js"
API_RESULTS_FILE = ROOT / "data" / "api-results.json"
FIXTURE_TIMEZONE = ZoneInfo("Europe/London")
PRIOR_SEASON_WEIGHT = 1.0
CURRENT_SEASON_WEIGHT = 3.0
NEUTRAL_PRIOR_MATCHES = 4.0


def parse_excel_datetime(value):
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, (int, float)):
        parsed = from_excel(value)
    else:
        raise ValueError(f"Unsupported Excel date value: {value!r}")

    if parsed.tzinfo:
        return parsed.astimezone(timezone.utc)
    return parsed.replace(tzinfo=FIXTURE_TIMEZONE).astimezone(timezone.utc)


def parse_result(value):
    if value is None or str(value).strip() == "":
        return None
    left, right = str(value).split("-")
    return int(left.strip()), int(right.strip())


def load_api_results():
    if not API_RESULTS_FILE.exists():
        return {}

    payload = json.loads(API_RESULTS_FILE.read_text(encoding="utf-8"))
    results = {}
    for item in payload.get("results", []):
        score = item.get("score") or {}
        home_score = score.get("home")
        away_score = score.get("away")
        if home_score is None or away_score is None:
            continue
        results[item["id"]] = {
            "result": (int(home_score), int(away_score)),
            "source": item.get("source") or "api",
        }
    return results


def source_path(path):
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def read_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for raw in ws.iter_rows(values_only=True):
        if raw[0] == "Match Number" or raw[0] is None:
            continue
        rows.append(
            {
                "match_number": int(raw[0]),
                "round": int(raw[1]),
                "date": parse_excel_datetime(raw[2]),
                "location": raw[3],
                "home": raw[4],
                "away": raw[5],
                "result": parse_result(raw[6]),
            }
        )
    return rows


def poisson_pmf(lam, goals):
    return math.exp(-lam) * (lam**goals) / math.factorial(goals)


def build_strengths(history_rows, current_rows):
    prior_played = [row for row in history_rows if row["result"]]
    current_played = [row for row in current_rows if row["result"]]
    weighted_rows = [
        *((row, PRIOR_SEASON_WEIGHT) for row in prior_played),
        *((row, CURRENT_SEASON_WEIGHT) for row in current_played),
    ]
    total_weight = sum(weight for _, weight in weighted_rows)
    if not total_weight:
        raise ValueError("At least one completed match is required to calculate team strengths.")

    total_home_goals = sum(row["result"][0] * weight for row, weight in weighted_rows)
    total_away_goals = sum(row["result"][1] * weight for row, weight in weighted_rows)
    league_home_avg = total_home_goals / total_weight
    league_away_avg = total_away_goals / total_weight

    buckets = defaultdict(lambda: defaultdict(int))
    teams = set()
    for row, weight in weighted_rows:
        home_score, away_score = row["result"]
        teams.add(row["home"])
        teams.add(row["away"])
        buckets[row["home"]]["home_played"] += weight
        buckets[row["home"]]["home_goals_for"] += home_score * weight
        buckets[row["home"]]["home_goals_against"] += away_score * weight
        buckets[row["away"]]["away_played"] += weight
        buckets[row["away"]]["away_goals_for"] += away_score * weight
        buckets[row["away"]]["away_goals_against"] += home_score * weight

    strengths = {}
    for team in teams:
        stats = buckets[team]
        home_played = stats["home_played"]
        away_played = stats["away_played"]
        strengths[team] = {
            "home_attack": (
                (stats["home_goals_for"] + NEUTRAL_PRIOR_MATCHES * league_home_avg)
                / (home_played + NEUTRAL_PRIOR_MATCHES)
            )
            / league_home_avg,
            "home_defence": (
                (stats["home_goals_against"] + NEUTRAL_PRIOR_MATCHES * league_away_avg)
                / (home_played + NEUTRAL_PRIOR_MATCHES)
            )
            / league_away_avg,
            "away_attack": (
                (stats["away_goals_for"] + NEUTRAL_PRIOR_MATCHES * league_away_avg)
                / (away_played + NEUTRAL_PRIOR_MATCHES)
            )
            / league_away_avg,
            "away_defence": (
                (stats["away_goals_against"] + NEUTRAL_PRIOR_MATCHES * league_home_avg)
                / (away_played + NEUTRAL_PRIOR_MATCHES)
            )
            / league_home_avg,
            "prior_source": "weighted-completed-history",
        }

    return {
        "league_home_avg": league_home_avg,
        "league_away_avg": league_away_avg,
        "prior_matches": len(prior_played),
        "current_matches": len(current_played),
        "teams": strengths,
    }


def team_strength(strengths, team):
    if team in strengths["teams"]:
        return strengths["teams"][team]
    return {
        "home_attack": 1.0,
        "home_defence": 1.0,
        "away_attack": 1.0,
        "away_defence": 1.0,
        "prior_source": "neutral-promoted-prior",
    }


def model_fixture(strengths, row, max_goals=8):
    home = team_strength(strengths, row["home"])
    away = team_strength(strengths, row["away"])

    home_lambda = strengths["league_home_avg"] * home["home_attack"] * away["away_defence"]
    away_lambda = strengths["league_away_avg"] * away["away_attack"] * home["home_defence"]
    home_goal_probs = [poisson_pmf(home_lambda, goals) for goals in range(max_goals + 1)]
    away_goal_probs = [poisson_pmf(away_lambda, goals) for goals in range(max_goals + 1)]

    home_win = draw = away_win = 0.0
    for home_goals, home_prob in enumerate(home_goal_probs):
        for away_goals, away_prob in enumerate(away_goal_probs):
            score_prob = home_prob * away_prob
            if home_goals > away_goals:
                home_win += score_prob
            elif home_goals == away_goals:
                draw += score_prob
            else:
                away_win += score_prob

    total = home_win + draw + away_win
    home_win /= total
    draw /= total
    away_win /= total
    home_decisive = home_win / (home_win + away_win)
    away_decisive = away_win / (home_win + away_win)

    favourite_decisive = max(home_decisive, away_decisive)
    if favourite_decisive < 0.60:
        favourite_pints, underdog_pints = 1, 1
    elif favourite_decisive < 0.715:
        favourite_pints, underdog_pints = 1, 2
    else:
        favourite_pints, underdog_pints = 1, 3

    if home_decisive >= away_decisive:
        home_pints, away_pints = favourite_pints, underdog_pints
    else:
        home_pints, away_pints = underdog_pints, favourite_pints

    return {
        "home_lambda": round(home_lambda, 4),
        "away_lambda": round(away_lambda, 4),
        "home_probability": round(home_win, 6),
        "draw_probability": round(draw, 6),
        "away_probability": round(away_win, 6),
        "home_decisive_probability": round(home_decisive, 6),
        "away_decisive_probability": round(away_decisive, 6),
        "home_fair_decimal_odds": round(1 / home_win, 3),
        "draw_fair_decimal_odds": round(1 / draw, 3),
        "away_fair_decimal_odds": round(1 / away_win, 3),
        "home_pints": home_pints,
        "away_pints": away_pints,
        "home_prior_source": home["prior_source"],
        "away_prior_source": away["prior_source"],
    }


def first_picker_for_round(rows_in_round):
    starting = "A" if rows_in_round[0]["round"] % 2 else "M"
    ordered = sorted(rows_in_round, key=lambda item: (item["date"], item["match_number"]))
    return {
        row["match_number"]: ("A" if index % 2 == 0 else "M") if starting == "A" else ("M" if index % 2 == 0 else "A")
        for index, row in enumerate(ordered)
    }


def select_upcoming_gameweek(fixture_rows, now):
    all_upcoming = [row for row in fixture_rows if row["result"] is None and row["date"] >= now]
    if not all_upcoming:
        return None, []

    gameweek = min(all_upcoming, key=lambda item: (item["date"], item["match_number"]))["round"]
    return gameweek, [row for row in all_upcoming if row["round"] == gameweek]


def load_previous_payload():
    if not OUT_FILE.exists():
        return {}

    text = OUT_FILE.read_text(encoding="utf-8")
    prefix = "window.AVSM_DATA = "
    if not text.startswith(prefix):
        return {}
    return json.loads(text.removeprefix(prefix).removesuffix(";\n"))


def saved_prediction(saved, fallback):
    if not saved:
        return fallback

    field_map = {
        "home_pints": "homePints",
        "away_pints": "awayPints",
        "home_probability": "homeProbability",
        "draw_probability": "drawProbability",
        "away_probability": "awayProbability",
        "home_fair_decimal_odds": "homeOdds",
        "draw_fair_decimal_odds": "drawOdds",
        "away_fair_decimal_odds": "awayOdds",
    }
    prediction = dict(fallback)
    for prediction_key, saved_key in field_map.items():
        if saved.get(saved_key) is not None:
            prediction[prediction_key] = saved[saved_key]
    return prediction


def build_payload(now, preserve_odds=False):
    history_rows = read_rows(HISTORY_FILE)
    fixture_rows = read_rows(FIXTURE_FILE)
    api_results = load_api_results()
    previous_payload = load_previous_payload()
    previous_matches = {
        item["id"]: item
        for item in [
            *previous_payload.get("fixtures", []),
            *previous_payload.get("recordedResults", []),
        ]
    }

    for row in fixture_rows:
        api_result = api_results.get(f"epl2026-{row['match_number']}")
        if api_result:
            row["result"] = api_result["result"]
            row["result_source"] = api_result["source"]
        elif row["result"]:
            row["result_source"] = "excel"
        else:
            row["result_source"] = None

    completed_current_rows = [row for row in fixture_rows if row["result"] and row["date"] < now]
    strengths = build_strengths(history_rows, completed_current_rows)

    first_pickers = {}
    rounds = defaultdict(list)
    for row in fixture_rows:
        rounds[row["round"]].append(row)
    for rows in rounds.values():
        first_pickers.update(first_picker_for_round(rows))

    past = [row for row in fixture_rows if row["date"] < now]
    upcoming_gameweek, upcoming = select_upcoming_gameweek(fixture_rows, now)

    previous_gameweek = previous_payload.get("upcomingGameweek")
    preserving_current_week = preserve_odds and previous_gameweek == upcoming_gameweek
    odds_updated_at = (
        previous_payload.get("oddsUpdatedAt")
        if preserving_current_week and previous_payload.get("oddsUpdatedAt")
        else now.isoformat().replace("+00:00", "Z")
    )

    fixtures = []
    for row in upcoming:
        prediction = model_fixture(strengths, row)
        fixture_id = f"epl2026-{row['match_number']}"
        saved = previous_matches.get(fixture_id) if preserving_current_week else None
        prediction = saved_prediction(saved, prediction)
        fixtures.append(
            {
                "id": fixture_id,
                "season": "2026/27",
                "gameweek": row["round"],
                "sequence": row["match_number"],
                "kickoff": row["date"].isoformat().replace("+00:00", "Z"),
                "location": row["location"],
                "home": row["home"],
                "away": row["away"],
                "firstPicker": first_pickers[row["match_number"]],
                "homePints": prediction["home_pints"],
                "awayPints": prediction["away_pints"],
                "homeProbability": prediction["home_probability"],
                "drawProbability": prediction["draw_probability"],
                "awayProbability": prediction["away_probability"],
                "homeOdds": prediction["home_fair_decimal_odds"],
                "drawOdds": prediction["draw_fair_decimal_odds"],
                "awayOdds": prediction["away_fair_decimal_odds"],
                "status": "open",
                "score": None,
                "pintsLockedAt": saved.get("pintsLockedAt") if saved else odds_updated_at,
            }
        )

    history = []
    for row in past:
        prediction = model_fixture(strengths, row)
        fixture_id = f"epl2026-{row['match_number']}"
        saved = previous_matches.get(fixture_id)
        prediction = saved_prediction(saved, prediction)
        score = None
        status = "awaiting"
        if row["result"]:
            home_score, away_score = row["result"]
            score = {"home": home_score, "away": away_score}
            status = "recorded"

        history.append(
            {
                "id": fixture_id,
                "season": "2026/27",
                "gameweek": row["round"],
                "sequence": row["match_number"],
                "kickoff": row["date"].isoformat().replace("+00:00", "Z"),
                "location": row["location"],
                "home": row["home"],
                "away": row["away"],
                "firstPicker": first_pickers[row["match_number"]],
                "homePints": prediction["home_pints"],
                "awayPints": prediction["away_pints"],
                "homeProbability": prediction["home_probability"],
                "drawProbability": prediction["draw_probability"],
                "awayProbability": prediction["away_probability"],
                "homeOdds": prediction["home_fair_decimal_odds"],
                "drawOdds": prediction["draw_fair_decimal_odds"],
                "awayOdds": prediction["away_fair_decimal_odds"],
                "status": status,
                "score": score,
                "resultSource": row["result_source"],
                "pintsLockedAt": saved.get("pintsLockedAt") if saved else row["date"].isoformat().replace("+00:00", "Z"),
            }
        )

    return {
        "generatedAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "predictionStartsAt": now.isoformat().replace("+00:00", "Z"),
        "oddsUpdatedAt": odds_updated_at,
        "upcomingGameweek": upcoming_gameweek,
        "sourceFiles": {
            "strengths": f"{HISTORY_FILE.name} and completed 2026/27 results",
            "fixtures": FIXTURE_FILE.name,
            "fixtureTimezone": str(FIXTURE_TIMEZONE),
            "apiResults": source_path(API_RESULTS_FILE),
        },
        "model": {
            "version": "weighted-poisson-completed-history-v2",
            "leagueHomeGoals": round(strengths["league_home_avg"], 4),
            "leagueAwayGoals": round(strengths["league_away_avg"], 4),
            "priorSeasonMatches": strengths["prior_matches"],
            "currentSeasonMatches": strengths["current_matches"],
            "priorSeasonWeight": PRIOR_SEASON_WEIGHT,
            "currentSeasonWeight": CURRENT_SEASON_WEIGHT,
            "neutralPriorMatches": NEUTRAL_PRIOR_MATCHES,
            "maxGoals": 8,
            "pintThresholds": {
                "oneOneBelow": 0.60,
                "oneTwoBelow": 0.715,
            },
        },
        "fixtures": sorted(fixtures, key=lambda item: (item["kickoff"], item["sequence"])),
        "recordedResults": sorted(history, key=lambda item: (item["kickoff"], item["sequence"])),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--now", help="UTC timestamp to use as the prediction cutoff, e.g. 2026-09-09T03:04:18Z")
    parser.add_argument(
        "--preserve-odds",
        action="store_true",
        help="Keep the current gameweek's published odds while refreshing fixtures and results.",
    )
    args = parser.parse_args()
    if args.now:
        now = datetime.fromisoformat(args.now.replace("Z", "+00:00")).astimezone(timezone.utc)
    else:
        now = datetime.now(timezone.utc)

    payload = build_payload(now, preserve_odds=args.preserve_odds)
    OUT_FILE.write_text(
        "window.AVSM_DATA = "
        + json.dumps(payload, indent=2)
        + ";\n",
        encoding="utf-8",
    )
    print(f"Wrote {OUT_FILE}")
    print(f"Upcoming gameweek: {payload['upcomingGameweek']}")
    print(f"Upcoming fixtures: {len(payload['fixtures'])}")
    print(f"Recorded 2026/27 results before cutoff: {len(payload['recordedResults'])}")


if __name__ == "__main__":
    main()
